#!/usr/bin/env python3
"""Lagar data/program.json frå treningsprogram_v2.xlsx og ovingar_video.json.
Køyr: python3 scripts/build_data.py [--offentleg] [sti/til/xlsx] [sti/til/json]"""
import json, os, re, sys, datetime as dt, unicodedata
from pathlib import Path
import openpyxl

# Bruk: python3 scripts/build_data.py [--offentleg] [arbeidsbok.xlsx] [ovingar_video.json]
# --offentleg: for den publiserte versjonen (GitHub Pages) – fjernar startvekt og brukar private
#   erstatningsreglar frå /workspace/build/offentleg_erstatt.json (eller $TRENING_OFFENTLEG_REGLAR).
OFFENTLEG = "--offentleg" in sys.argv
ARG = [a for a in sys.argv[1:] if not a.startswith("--")]
XLSX = Path(ARG[0] if len(ARG) > 0 else "/workspace/treningsprogram_v2.xlsx")
VJSON = Path(ARG[1] if len(ARG) > 1 else "/workspace/ovingar_video.json")
OUT = Path(__file__).resolve().parent.parent / "data" / "program.json"

wb = openpyxl.load_workbook(XLSX)
vid = json.loads(VJSON.read_text(encoding="utf-8"))["ovingar"]
by_prog = {}
for v in vid:
    for n in v["programnamn"]:
        by_prog[n] = v["id"]

def slug(s):
    s = s.replace("ø", "o").replace("æ", "ae").replace("å", "a").replace("Ø", "O").replace("Æ", "Ae").replace("Å", "A")
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")

def parse_sets(txt):
    """'3 × 8–12 per side' -> dict"""
    t = txt.replace("–", "-")
    m = re.match(r"\s*(\d+)(?:-(\d+))?\s*×\s*(\d+)(?:-(\d+))?\s*(m|s)?\b", t)
    if not m:
        return {"sets": 1, "repMin": None, "repMax": None, "unit": "rep", "perSide": "per side" in t}
    s_lo, s_hi, r_lo, r_hi, unit = m.groups()
    return {"sets": int(s_hi or s_lo), "setsMin": int(s_lo), "repMin": int(r_lo), "repMax": int(r_hi or r_lo),
            "unit": {"m": "m", "s": "s"}.get(unit, "rep"), "perSide": "per side" in t}

def parse_kg(txt):
    t = (txt or "").strip()
    low = t.lower()
    if "kg" in low:
        nums = re.findall(r"(\d+(?:[.,]\d+)?)\s*(?:–\s*\d+\s*)?kg", t)
        # «2 × 12 kg» -> 12 ; «16 kg → 18 kg» -> 16
        kg = float(nums[0].replace(",", ".")) if nums else None
        two = bool(re.search(r"\b2\s*×\s*\d", t)) or "per hand" in low
        start = kg
        if low.startswith("utan vekt"):
            start = 0.0
        return {"type": "manual", "startKg": start, "targetKg": kg, "twoDumbbells": two}
    if "strikk" in low:
        return {"type": "strikk", "startKg": None}
    if "kroppsvekt" in low:
        return {"type": "kroppsvekt", "startKg": None}
    return {"type": "ingen", "startKg": None}

def parse_rest(txt):
    nums = [int(n) for n in re.findall(r"\d+", txt or "")]
    return max(nums) if nums else 0

def header_map(ws, row=3):
    return {c.value: c.column for c in ws[row] if c.value}

def program(sheet, person):
    ws = wb[sheet]
    h = header_map(ws)
    def g(r, name):
        col = h.get(name)
        return ws.cell(row=r, column=col).value if col else None
    ex, notes = [], []
    for r in range(4, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if a in ("A", "B"):
            name = g(r, "Øving")
            kg = parse_kg(g(r, "Start-kg (manual 3–24 kg)"))
            sets = parse_sets(g(r, "Sett × rep"))
            vkey = by_prog.get(name)
            note = g(r, "Ryggmerknad") if person == "atle" else g(r, "Merknad")
            ex.append({
                "id": f"{person}-{a}-{g(r,'Nr')}-{slug(name)}",
                "okt": a, "nr": g(r, "Nr"), "namn": name,
                "gruppe": g(r, "Gruppe"),
                "status": g(r, "Frå v1 / status"),
                "korleis": g(r, "Korleis (kort)"),
                "settRep": g(r, "Sett × rep"), **sets,
                "startKgTekst": g(r, "Start-kg (manual 3–24 kg)"), **kg,
                "lett": g(r, "Lett-variant (urolig rygg)") or g(r, "Lett-variant"),
                "pauseTekst": g(r, "Pause"), "pauseSek": parse_rest(g(r, "Pause")),
                "merknad": note or "",
                "atvaring": "⚠" in (note or "") or "⚠" in (g(r, "Frå v1 / status") or ""),
                "video": vkey,
            })
            assert vkey, f"manglar video for {name}"
        elif isinstance(a, str) and r > 5 and not a.startswith("ØKT") and len(a) > 40:
            notes.append(a)
    return {"ovingar": ex, "notat": notes}

def start_her():
    ws = wb["Start her"]
    vals = [(r, ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value) for r in range(1, ws.max_row + 1)]
    out = {"oppvarming": None, "reglar": [], "utstyr": [], "atvaring": None, "ingress": ws.cell(row=2, column=1).value}
    sec = None
    for r, a, b in vals:
        if isinstance(a, str) and a.isupper() and len(a) > 5:
            sec = a; continue
        if not a: continue
        if sec and sec.startswith("OPPVARMING"):
            out["oppvarming"] = a; sec = None
        elif sec and sec.startswith("REGLAR") and a != "Regel":
            out["reglar"].append({"tittel": a, "tekst": b})
        elif sec and sec.startswith("UTSTYR") and a != "Utstyr":
            out["utstyr"].append({"utstyr": a, "tekst": b})
        elif isinstance(a, str) and a.startswith("Dette er eit generelt"):
            out["atvaring"] = a
    return out

def kalender():
    ws = wb["Vekekalender"]; h = header_map(ws)
    days = []
    for r in range(4, 32):
        d = ws.cell(row=r, column=h["Dato"]).value
        if not isinstance(d, (dt.date, dt.datetime)): continue
        days.append({"dato": d.strftime("%Y-%m-%d"), "vekedag": ws.cell(row=r, column=h["Vekedag"]).value,
                     "veke": ws.cell(row=r, column=h["Veke"]).value, "type": ws.cell(row=r, column=h["Type"]).value,
                     "okt": ws.cell(row=r, column=h["Økt"]).value or None,
                     "atle": ws.cell(row=r, column=h["Atle"]).value, "audny": ws.cell(row=r, column=h["Audny"]).value,
                     "tid": ws.cell(row=r, column=h["Tid"]).value})
    n = 0
    for d in days:
        if d["type"] == "Styrke":
            n += 1; d["oktNr"] = n
    return days

def progresjon():
    ws = wb["Progresjon"]
    return [{"steg": ws.cell(row=r, column=1).value, "regel": ws.cell(row=r, column=2).value, "doeme": ws.cell(row=r, column=3).value}
            for r in range(4, ws.max_row + 1) if ws.cell(row=r, column=1).value]

def framgang():
    ws = wb["Framgang"]
    res = {}
    for r in range(5, 12):
        lab = ws.cell(row=r, column=1).value
        if lab:
            res[lab] = {"atle": ws.cell(row=r, column=2).value, "audny": ws.cell(row=r, column=3).value}
    return res

def guide():
    ws = wb["Øvingsguide"]; h = header_map(ws)
    vmap = {v["ovingsnamn"]: v for v in vid}
    out = []
    for r in range(4, ws.max_row + 1):
        name = ws.cell(row=r, column=h["Øving"]).value
        if name not in vmap: continue
        v = vmap[name]
        out.append({"id": v["id"], "namn": name, "gjeld": ws.cell(row=r, column=h["Gjeld"]).value,
                    "teknikk": ws.cell(row=r, column=h["Slik ser god teknikk ut"]).value,
                    "feil": ws.cell(row=r, column=h["Typiske feil"]).value,
                    "videoUrl": v["video_url"], "videoTittel": v["video_tittel"], "kjelde": v["kjelde"],
                    "reserveUrl": v["reserve_url"], "reserveTittel": v["reserve_tittel"], "reserveKjelde": v["reserve_kjelde"],
                    "merknad": v["merknad"]})
    assert len(out) == len(vid), (len(out), len(vid))
    return out

fg = framgang()
data = {
    "laga": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
    "kjelde": XLSX.name,
    "blokk": {"start": "2026-10-19", "slutt": "2026-11-15"},
    "manual": {"min": 3, "max": 24},
    "personar": {
        "atle": {"namn": "Atle", "scoreNamn": "Rygg", "scoreTekst": "Rygg 1–5: 1 = ingen plage, 3 = kjennbar men styrbar, 5 = skarp/utstrålande.",
                 "malVekt": fg.get("Mål", {}).get("atle"), "startVekt": fg.get("Startvekt v1 (21.09.2026)", {}).get("atle"),
                 **program("Atle – økter", "atle")},
        "audny": {"namn": "Audny", "scoreNamn": "Nakke/skulder", "scoreTekst": "Nakke/skulder 1–5: 1 = ingen plage, 3 = kjennbar men styrbar, 5 = skarp.",
                  "malVekt": fg.get("Mål", {}).get("audny"), "minVekt": 49, "startVekt": fg.get("Startvekt v1 (21.09.2026)", {}).get("audny"),
                  **program("Audny – økter", "audny")},
    },
    "start": start_her(),
    "kalender": kalender(),
    "progresjon": progresjon(),
    "guide": guide(),
}
# Tekst som viser til Excel-arka blir skriven om for appen
def fix(o):
    if isinstance(o, str):
        for a, b in [("i arket «Øvingsguide»", "under «Øvingar»"), ("(sjå Start her)", "(sjå «Oppvarming» øvst i økta)"),
                     ("Sjå arket «Øvingsguide» for reservevideo og lett-variantar.", "Reservevideo finn du under «Øvingar»."),
                     ("sjå «Progresjon»", "sjå «Meir → Progresjon»"), ("(sjå arket Progresjon)", "(sjå «Progresjon» under «Meir»)"), ("Fyll inn det du faktisk brukar i «Logg Atle».", "Logg det du faktisk brukar.")]:
            o = o.replace(a, b)
        return o
    if isinstance(o, list): return [fix(x) for x in o]
    if isinstance(o, dict): return {k: fix(v) for k, v in o.items()}
    return o
data = fix(data)
if OFFENTLEG:
    for pp in data["personar"].values():
        pp.pop("startVekt", None)
    data.pop("kjelde", None)
    tekst = json.dumps(data, ensure_ascii=False)
    # Private erstatningsreglar (ligg IKKJE i repoet): {"erstatt": [[frå, til], ...], "forbode": [ord, ...]}
    regelfil = Path(os.environ.get("TRENING_OFFENTLEG_REGLAR", "/workspace/build/offentleg_erstatt.json"))
    if not regelfil.exists():
        sys.exit(f"Manglar {regelfil} – trengst for --offentleg")
    reglar = json.loads(regelfil.read_text(encoding="utf-8"))
    for a, b in reglar.get("erstatt", []):
        tekst = tekst.replace(a, b)
    for ord_ in reglar.get("forbode", []):
        assert ord_.lower() not in tekst.lower(), f"forbode ord att i offentleg data ({regelfil})"
    data = json.loads(tekst)
OUT.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
print("skreiv", OUT, {p: len(v["ovingar"]) for p, v in data["personar"].items()}, "guide", len(data["guide"]), "dagar", len(data["kalender"]))
